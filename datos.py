"""
Carga de datos reales desde Excel.

Este archivo permite reemplazar los datos sintéticos por información real de una
empresa, exportada por ejemplo desde SAP y resumida por mes en una plantilla de Excel.

IMPORTANTE (gobierno de datos): use datos reales solo con autorización escrita de la
empresa, y preferiblemente ejecutando el aplicativo en un computador de la empresa,
no en la versión publicada en internet.
"""

import io
import unicodedata
from dataclasses import replace

import numpy as np
import pandas as pd

import modelo as m

# ---------------------------------------------------------------------------
# Estructura de la plantilla
# ---------------------------------------------------------------------------
# columna: (descripción, de dónde sale normalmente en SAP, obligatoria)
COLUMNAS = {
    "mes": ("Mes del dato, en formato AAAA-MM (por ejemplo 2025-01)", "", True),
    "facturado": ("Total facturado del mes por el servicio de aseo", "Facturación de ventas / FBL5N", True),
    "recaudo": ("Dinero efectivamente recaudado en el mes", "Pagos de clientes compensados / FBL5N", True),
    "disposicion_final": ("Pagos por disposición final de residuos", "FBL3N o KSB1 (cuenta o centro de costo)", False),
    "nomina": ("Pagos de nómina operativa", "FBL3N o KSB1", False),
    "combustible": ("Pagos de combustible", "FBL3N o KSB1", False),
    "mantenimiento": ("Pagos de mantenimiento de flota", "FBL3N o KSB1", False),
    "administrativos": ("Gastos administrativos pagados", "FBL3N o KSB1", False),
    "otros": ("Otros pagos operativos", "FBL3N", False),
    "poliza": ("Pagos de pólizas (registrar en el mes en que se pagan)", "FBL1N / FBL3N", False),
    "servicio_deuda": ("Cuotas de créditos: capital e intereses", "FBL3N (obligaciones financieras)", False),
    "inversion_flota": ("Compra o reposición de vehículos y contenedores", "Activos fijos / FBL3N", False),
    "impuestos": ("Pagos de impuesto de renta", "FBL3N (impuestos por pagar)", False),
    "saldo_caja": ("Saldo de caja y bancos al cierre del mes", "FBL3N de cuentas de bancos / FF7A", True),
}
OBLIGATORIAS = [c for c, v in COLUMNAS.items() if v[2]]
EGRESOS = [c for c in COLUMNAS if c not in ("mes", "facturado", "recaudo", "saldo_caja")]
COLUMNAS_PRESUPUESTO = ["mes", "facturado", "recaudo"] + EGRESOS
MIN_MESES = 12
MESES_BACKTESTING = 24


# ---------------------------------------------------------------------------
# Plantilla de Excel para descargar
# ---------------------------------------------------------------------------
def _ejemplo_historico() -> pd.DataFrame:
    """36 meses de ejemplo (sintéticos) para que la plantilla se pueda probar de inmediato."""
    h = m.historia_sintetica(m.Supuestos()).copy()
    fechas = pd.date_range("2023-01-01", periods=len(h), freq="MS")
    h["mes"] = fechas.strftime("%Y-%m")
    return h[list(COLUMNAS)].round(1)


def _ejemplo_presupuesto() -> pd.DataFrame:
    p = m.proyectar(m.Supuestos(), "Base", meses=36, anio_inicio=2023,
                    caja_inicial=m.CAJA_INICIAL_HISTORIA).iloc[24:36].copy()
    p["mes"] = pd.date_range("2025-01-01", periods=12, freq="MS").strftime("%Y-%m")
    return p[COLUMNAS_PRESUPUESTO].round(1)


def plantilla_excel() -> bytes:
    from openpyxl.styles import Alignment, Font, PatternFill
    verde = PatternFill("solid", fgColor="14432A")
    claro = PatternFill("solid", fgColor="E8F1EA")
    blanco = Font(bold=True, color="FFFFFF")

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        inst = [
            ["PLANTILLA DE DATOS FINASEO"],
            [""],
            ["1. Todas las cifras van en MILLONES de pesos colombianos (por ejemplo, 11.250 = $11.250 millones)."],
            ["2. Una fila por mes. Mínimo 12 meses; con 24 o más se puede hacer el backtesting con datos reales."],
            ["3. La hoja 'Historico' es obligatoria. La hoja 'Presupuesto' es opcional (para comparar presupuesto y ejecutado)."],
            ["4. Los datos de ejemplo son SINTÉTICOS: reemplácelos por los de la empresa."],
            ["5. No cambie los nombres de las columnas. Si un rubro no aplica, déjelo en 0."],
            ["6. CONFIDENCIALIDAD: use datos reales solo con autorización escrita de la empresa, y preferiblemente "
             "en un computador de la empresa, no en la versión publicada en internet."],
            [""],
            ["Columna", "Qué debe contener", "Fuente habitual en SAP (verificar con sistemas)", "¿Obligatoria?"],
        ]
        for c, (desc, sap, obl) in COLUMNAS.items():
            inst.append([c, desc, sap, "Sí" if obl else "No (si falta, se toma como 0)"])
        pd.DataFrame(inst).to_excel(xw, sheet_name="Instrucciones", index=False, header=False)
        _ejemplo_historico().to_excel(xw, sheet_name="Historico", index=False)
        _ejemplo_presupuesto().to_excel(xw, sheet_name="Presupuesto", index=False)

        ws = xw.sheets["Instrucciones"]
        ws["A1"].font = Font(bold=True, size=14, color="14432A")
        ws["A8"].font = Font(bold=True, color="B42318")
        for cell in ws[10]:
            cell.fill, cell.font = verde, blanco
        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 62
        ws.column_dimensions["C"].width = 44
        ws.column_dimensions["D"].width = 30
        for hoja in ("Historico", "Presupuesto"):
            w = xw.sheets[hoja]
            for cell in w[1]:
                cell.fill, cell.font = verde, blanco
                cell.alignment = Alignment(horizontal="center")
            for col in w.iter_cols(min_col=2, min_row=2):
                for cell in col:
                    cell.number_format = "#,##0.0"
            for i, col in enumerate(w.columns, start=1):
                w.column_dimensions[col[0].column_letter].width = 11 if i == 1 else 16
            for r in range(2, w.max_row + 1, 2):
                for cell in w[r]:
                    cell.fill = claro
            w.freeze_panes = "B2"
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Lectura y validación
# ---------------------------------------------------------------------------
def _normalizar(texto) -> str:
    t = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return t.strip().lower().replace(" ", "_")


def _buscar_hoja(hojas: dict, nombre: str):
    for k, v in hojas.items():
        if _normalizar(k) == nombre:
            return v
    return None


def _preparar(df: pd.DataFrame, hoja: str, columnas: list[str], obligatorias: list[str],
              errores: list, avisos: list) -> pd.DataFrame | None:
    df = df.copy()
    df.columns = [_normalizar(c) for c in df.columns]
    df = df.dropna(how="all")
    faltan = [c for c in obligatorias if c not in df.columns]
    if faltan:
        errores.append(f"En la hoja '{hoja}' faltan columnas obligatorias: {', '.join(faltan)}.")
        return None
    for c in columnas:
        if c not in df.columns:
            df[c] = 0.0
            avisos.append(f"La hoja '{hoja}' no tiene la columna '{c}'; se tomó como 0.")
    df = df[columnas]

    fechas = pd.to_datetime(df["mes"].astype(str).str.strip().str[:7], format="%Y-%m", errors="coerce")
    malas = df.loc[fechas.isna(), "mes"].tolist()
    if malas:
        errores.append(f"En la hoja '{hoja}' hay meses con formato inválido: {malas[:5]}. Use AAAA-MM, por ejemplo 2025-01.")
        return None
    df["fecha"] = fechas
    if df["fecha"].duplicated().any():
        errores.append(f"En la hoja '{hoja}' hay meses repetidos: "
                       f"{df.loc[df['fecha'].duplicated(), 'fecha'].dt.strftime('%Y-%m').tolist()[:5]}.")
        return None
    df = df.sort_values("fecha").reset_index(drop=True)

    numericas = [c for c in columnas if c != "mes"]
    for c in numericas:
        convertido = pd.to_numeric(df[c], errors="coerce")
        filas = df.loc[convertido.isna() & df[c].notna(), "fecha"].dt.strftime("%Y-%m").tolist()
        if filas:
            errores.append(f"En la hoja '{hoja}', la columna '{c}' tiene valores que no son números en: {filas[:5]}.")
        df[c] = convertido.fillna(0.0)
    if errores:
        return None

    negativos = [c for c in EGRESOS + ["facturado", "recaudo"] if c in df and (df[c] < 0).any()]
    if negativos:
        avisos.append(f"En la hoja '{hoja}' hay valores negativos en: {', '.join(negativos)}. "
                      "Los pagos deben registrarse como números positivos.")
    if df["facturado"].mean() > 1_000_000:
        errores.append(f"En la hoja '{hoja}' las cifras parecen estar en pesos y no en millones "
                       "(la facturación promedio supera el billón). Divida los valores entre 1.000.000.")
        return None

    esperado = pd.date_range(df["fecha"].min(), df["fecha"].max(), freq="MS")
    if len(esperado) != len(df):
        avisos.append(f"En la hoja '{hoja}' faltan meses intermedios; revise que la serie sea continua.")
    df["mes"] = [f"{m.MESES_ES[f.month - 1]} {f.year}" for f in df["fecha"]]
    return df


def leer_excel(archivo) -> tuple[pd.DataFrame | None, pd.DataFrame | None, list[str], list[str]]:
    """Devuelve (historico, presupuesto, errores, avisos)."""
    errores, avisos = [], []
    try:
        hojas = pd.read_excel(archivo, sheet_name=None)
    except Exception as e:
        return None, None, [f"No se pudo leer el archivo como Excel: {e}"], avisos

    h = _buscar_hoja(hojas, "historico")
    if h is None:
        return None, None, ["El archivo no tiene una hoja llamada 'Historico'. Use la plantilla."], avisos
    hist = _preparar(h, "Historico", list(COLUMNAS), OBLIGATORIAS, errores, avisos)
    if hist is None:
        return None, None, errores, avisos
    if len(hist) < MIN_MESES:
        return None, None, [f"La hoja 'Historico' tiene {len(hist)} meses; se necesitan al menos {MIN_MESES}."], avisos
    if len(hist) < MESES_BACKTESTING:
        avisos.append(f"Con {len(hist)} meses no es posible el backtesting con datos reales "
                      f"(se necesitan {MESES_BACKTESTING}); se mantendrá el backtesting sintético.")

    hist["egresos"] = hist[EGRESOS].sum(axis=1)
    hist["flujo_neto"] = hist["recaudo"] - hist["egresos"]
    dif = (hist["saldo_caja"].diff() - hist["flujo_neto"]).iloc[1:].abs()
    if (dif > 0.05 * hist["facturado"].mean()).mean() > 0.5:
        avisos.append("En la mayoría de los meses la variación del saldo de caja no coincide con "
                      "recaudo menos egresos. Puede haber ingresos o pagos que no están en la plantilla "
                      "(por ejemplo, créditos nuevos o rendimientos); el modelo solo proyecta los rubros incluidos.")

    pres = None
    p = _buscar_hoja(hojas, "presupuesto")
    if p is not None and len(p.dropna(how="all")):
        err_p = []
        pres = _preparar(p, "Presupuesto", COLUMNAS_PRESUPUESTO, ["mes", "recaudo"], err_p, avisos)
        if err_p:
            avisos.extend(err_p + ["La hoja 'Presupuesto' se ignoró por los errores anteriores."])
            pres = None
        elif pres is not None:
            pres["egresos"] = pres[EGRESOS].sum(axis=1)
            pres["flujo_neto"] = pres["recaudo"] - pres["egresos"]
            comunes = set(pres["fecha"]) & set(hist["fecha"])
            if not comunes:
                avisos.append("Los meses de la hoja 'Presupuesto' no coinciden con los del histórico; "
                              "no se podrá comparar presupuesto y ejecutado.")
                pres = None
    return hist, pres, errores, avisos


# ---------------------------------------------------------------------------
# Calibración de los supuestos con datos reales
# ---------------------------------------------------------------------------
def calibrar(hist: pd.DataFrame, base: m.Supuestos, usuarios: int, toneladas: float) -> tuple[m.Supuestos, pd.DataFrame]:
    """Estima los drivers del modelo a partir del histórico real.
    Devuelve los supuestos calibrados y una tabla que explica cada cálculo."""
    ult = hist.tail(12)
    ult3 = hist.tail(3)
    rec = hist["recaudo"].iloc[1:].sum()
    fac_prev = hist["facturado"].shift(1).iloc[1:].sum()
    tasa = float(np.clip(rec / fac_prev, 0.5, 0.995)) if fac_prev > 0 else base.tasa_recaudo

    crec = base.crecimiento_usuarios_anual
    nota_crec = "Se mantiene el supuesto: se necesitan 24 meses para estimarlo"
    if len(hist) >= 24:
        f_ult, f_prev = hist["facturado"].tail(12).sum(), hist["facturado"].iloc[-24:-12].sum()
        if f_prev > 0:
            crec = float(np.clip(f_ult / f_prev - 1 - base.inflacion_anual, -0.05, 0.10))
            nota_crec = "Crecimiento anual de la facturación menos la inflación"

    usuarios = max(int(usuarios), 1)
    toneladas = max(float(toneladas), 1.0)
    s = replace(
        base,
        usuarios=usuarios,
        tarifa_promedio=float(ult3["facturado"].mean() / usuarios),
        crecimiento_usuarios_anual=crec,
        tasa_recaudo=tasa,
        toneladas_mes=toneladas,
        costo_disposicion_ton=float(ult["disposicion_final"].mean() / toneladas),
        nomina_operativa=float(ult["nomina"].mean()),
        combustible=float(ult["combustible"].mean()),
        mantenimiento_flota=float(ult["mantenimiento"].mean()),
        gastos_administrativos=float(ult["administrativos"].mean()),
        otros_pagos=float(ult["otros"].mean()),
        poliza_semestral=float(ult["poliza"].sum() / 2),
        servicio_deuda=float(ult["servicio_deuda"].mean()),
        inversion_flota=float(ult["inversion_flota"].mean()),
        impuesto_renta_anual=float(ult["impuestos"].sum()),
        caja_inicial=float(hist["saldo_caja"].iloc[-1]),
    )
    explicacion = pd.DataFrame([
        ["Usuarios facturados", f"{usuarios:,}".replace(",", "."), "Dato ingresado (SUI o área comercial)"],
        ["Tarifa promedio por usuario", f"${s.tarifa_promedio * 1e6:,.0f}".replace(",", "."), "Facturación promedio de los últimos 3 meses ÷ usuarios"],
        ["Crecimiento anual de usuarios", f"{crec * 100:.1f}%".replace(".", ","), nota_crec],
        ["Tasa de recaudo", f"{tasa * 100:.1f}%".replace(".", ","), "Recaudo total ÷ facturación del mes anterior"],
        ["Recaudo que llega al mes siguiente", f"{base.porcion_mes_siguiente * 100:.0f}%",
         "Se mantiene el supuesto: no se puede estimar con totales mensuales (requiere detalle por factura)"],
        ["Costo de disposición por tonelada", f"${s.costo_disposicion_ton * 1e6:,.0f}".replace(",", "."), "Promedio de 12 meses ÷ toneladas ingresadas"],
        ["Nómina, combustible, mantenimiento, administrativos, otros, deuda y flota", "Promedios mensuales", "Promedio de los últimos 12 meses"],
        ["Póliza semestral", f"${s.poliza_semestral:,.0f} M".replace(",", "."), "Pagos de pólizas de 12 meses ÷ 2 (el modelo las ubica en abril y octubre)"],
        ["Impuesto de renta anual", f"${s.impuesto_renta_anual:,.0f} M".replace(",", "."), "Pagos de 12 meses (el modelo los ubica en febrero, abril y junio)"],
        ["Caja inicial de la proyección", f"${s.caja_inicial:,.0f} M".replace(",", "."), f"Saldo de caja de {hist['mes'].iloc[-1]}"],
    ], columns=["Driver", "Valor calibrado", "Cómo se calculó"])
    return s, explicacion


def inicio_proyeccion(hist: pd.DataFrame) -> tuple[int, int]:
    """Año y mes del primer mes a proyectar: el siguiente al último dato real."""
    siguiente = hist["fecha"].iloc[-1] + pd.offsets.MonthBegin(1)
    return int(siguiente.year), int(siguiente.month)


# ---------------------------------------------------------------------------
# Backtesting con datos reales
# ---------------------------------------------------------------------------
def backtesting_real(hist: pd.DataFrame, base: m.Supuestos, usuarios: int, toneladas: float):
    """Usa todos los meses menos los últimos 12 para calibrar, pronostica esos 12 meses
    con el modelo y con el método tradicional, y compara contra lo realmente ocurrido."""
    entreno, prueba = hist.iloc[:-12], hist.iloc[-12:].reset_index(drop=True)
    s_ent, _ = calibrar(entreno, base, usuarios, toneladas)
    anio, mes = inicio_proyeccion(entreno)
    modelo = m.proyectar(s_ent, "Base", meses=12, anio_inicio=anio, mes_inicio=mes,
                         caja_inicial=float(entreno["saldo_caja"].iloc[-1]),
                         facturado_previo=entreno["facturado"].tail(2).tolist())
    tradicional = entreno["saldo_caja"].iloc[-1] + (entreno["flujo_neto"].tail(12).values
                                                   * (1 + base.inflacion_anual)).cumsum()
    comp = pd.DataFrame({"mes": prueba["mes"], "Real": prueba["saldo_caja"],
                         "Modelo por drivers": modelo["saldo_caja"].values,
                         "Método tradicional": tradicional})
    if (comp["Real"].abs() < 1e-9).any():
        res = {"mape_modelo": float("nan"), "mape_tradicional": float("nan")}
    else:
        res = {"mape_modelo": m.mape(comp["Real"], comp["Modelo por drivers"]),
               "mape_tradicional": m.mape(comp["Real"], comp["Método tradicional"])}
    return comp, res
